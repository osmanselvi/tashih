"""Düzeltilmiş metni dosya biçimlerine aktarma."""
from __future__ import annotations

import io
from collections.abc import Sequence


def _pages_text(pages: Sequence[tuple[int, str]]) -> str:
    return "\n\n".join(text for _, text in pages)


def _pages_markdown(pages: Sequence[tuple[int, str]]) -> str:
    parts = [f"## Sayfa {number}\n\n{text.strip()}" for number, text in pages]
    return "\n\n".join(parts) + "\n"


def _pages_docx(pages: Sequence[tuple[int, str]]) -> bytes:
    from docx import Document

    document = Document()
    for index, (number, text) in enumerate(pages):
        if index:
            document.add_page_break()
        document.add_heading(f"Sayfa {number}", level=2)
        for paragraph in text.split("\n"):
            if paragraph.strip():
                document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def export_document(pages: Sequence[tuple[int, str]], fmt: str) -> bytes:
    """(sayfa_no, metin) listesini istenen biçimde serileştirir."""
    if fmt == "txt":
        return (_pages_text(pages) + "\n").encode("utf-8")
    if fmt == "md":
        return _pages_markdown(pages).encode("utf-8")
    if fmt == "docx":
        return _pages_docx(pages)
    raise ValueError(f"Desteklenmeyen biçim: {fmt}")
